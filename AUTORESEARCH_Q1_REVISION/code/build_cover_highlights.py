"""Build Cover_letter / Highlights .docx and .pdf from their markdown sources."""
import re
import sys
from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from docx.shared import Pt
from reportlab.lib.enums import TA_JUSTIFY, TA_RIGHT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer

PKG = Path(__file__).resolve().parents[1] / "MATERIALS_DESIGN_TRANSFER"
MAX_HIGHLIGHT = 85


def parse(md_path):
    title, blocks = None, []
    for raw in md_path.read_text(encoding="utf-8").split("\n\n"):
        text = raw.strip()
        if not text:
            continue
        if text.startswith("# "):
            title = text[2:].strip()
        elif text.startswith("- "):
            blocks.append(("list", [l[2:].strip() for l in text.splitlines()]))
        else:
            blocks.append(("para", [l.rstrip() for l in text.splitlines()]))
    return title, blocks


def md_inline_to_rl(s):
    s = escape(s)
    return re.sub(r"\*(.+?)\*", r"<i>\1</i>", s)


def add_docx_runs(par, s):
    for i, chunk in enumerate(re.split(r"\*(.+?)\*", s)):
        if chunk:
            par.add_run(chunk).italic = bool(i % 2)


def build(name, md_file):
    title, blocks = parse(PKG / md_file)

    doc = Document()
    doc.styles["Normal"].font.name = "Times New Roman"
    doc.styles["Normal"].font.size = Pt(11)
    doc.add_heading(title, level=1)
    for kind, lines in blocks:
        if kind == "list":
            for item in lines:
                add_docx_runs(doc.add_paragraph(style="List Bullet"), item)
        else:
            p = doc.add_paragraph()
            for j, line in enumerate(lines):
                add_docx_runs(p, line)
                if j < len(lines) - 1:
                    p.add_run().add_break()
    doc.save(PKG / f"{name}.docx")

    ss = getSampleStyleSheet()
    body = ParagraphStyle("body", parent=ss["Normal"], fontName="Times-Roman",
                          fontSize=10.5, leading=13.5, alignment=TA_JUSTIFY, spaceAfter=7)
    story = [Paragraph(escape(title), ParagraphStyle("t", parent=ss["Title"], fontName="Times-Bold")),
             Spacer(1, 0.3 * cm)]
    for kind, lines in blocks:
        if kind == "list":
            story.append(ListFlowable(
                [ListItem(Paragraph(md_inline_to_rl(l), body)) for l in lines],
                bulletType="bullet", start="•"))
        elif re.fullmatch(r"\d{1,2} \w+ \d{4}", lines[0]):
            story.append(Paragraph(lines[0], ParagraphStyle("d", parent=body, alignment=TA_RIGHT)))
        else:
            story.append(Paragraph("<br/>".join(md_inline_to_rl(l) for l in lines), body))
    SimpleDocTemplate(str(PKG / f"{name}.pdf"), pagesize=A4, leftMargin=2.3 * cm,
                      rightMargin=2.3 * cm, topMargin=2.2 * cm, bottomMargin=2 * cm).build(story)
    return blocks


if __name__ == "__main__":
    build("Cover_letter", "COVER_LETTER.md")
    hl = build("Highlights", "HIGHLIGHTS.md")
    items = [l for kind, lines in hl if kind == "list" for l in lines]
    too_long = [(len(l), l) for l in items if len(l) > MAX_HIGHLIGHT]
    print(f"{len(items)} highlights, max {max(len(l) for l in items)} chars")
    if too_long:
        sys.exit(f"highlights over {MAX_HIGHLIGHT} chars: {too_long}")
